from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

# Configuration
INPUT_PRICE = 0.50       # USD per million input tokens
OUTPUT_PRICE = 1.50      # USD per million output tokens
USD_TO_INR = 88.0

BASELINE_INPUT = 1975
OPTIMISED_INPUT = 787
OUTPUT_TOKENS = 200

output_dir = Path("outputs")
output_dir.mkdir(exist_ok=True)

workbook = Workbook()

# Sheet 1: Cost comparison
ws = workbook.active
ws.title = "Cost Comparison"

ws.append(["LLM Cost Optimisation Audit"])
ws.append(["Metric", "Baseline", "Optimised", "Savings"])

rows = [
    ["Input Tokens", BASELINE_INPUT, OPTIMISED_INPUT, "=B3-C3"],
    ["Output Tokens", OUTPUT_TOKENS, OUTPUT_TOKENS, "=B4-C4"],
    ["Total Tokens", "=B3+B4", "=C3+C4", "=B5-C5"],
    ["Input Cost (USD)", f"=B3*{INPUT_PRICE}/1000000",
     f"=C3*{INPUT_PRICE}/1000000", "=B6-C6"],
    ["Output Cost (USD)", f"=B4*{OUTPUT_PRICE}/1000000",
     f"=C4*{OUTPUT_PRICE}/1000000", "=B7-C7"],
    ["Total Cost (USD)", "=B6+B7", "=C6+C7", "=B8-C8"],
    ["Total Cost (INR)", f"=B8*{USD_TO_INR}",
     f"=C8*{USD_TO_INR}", "=B9-C9"],
    ["Input Token Reduction (%)", 0,
     "=(B3-C3)/B3", "=C10"],
    ["Cost Reduction (%)", 0,
     "=(B9-C9)/B9", "=C11"],
]

for row in rows:
    ws.append(row)

# Sheet 2: Assumptions
assumptions = workbook.create_sheet("Assumptions")
assumption_rows = [
    ["Parameter", "Value"],
    ["Input price (USD / million tokens)", INPUT_PRICE],
    ["Output price (USD / million tokens)", OUTPUT_PRICE],
    ["USD to INR conversion rate", USD_TO_INR],
    ["Baseline input tokens", BASELINE_INPUT],
    ["Optimised input tokens", OPTIMISED_INPUT],
    ["Output tokens (held constant)", OUTPUT_TOKENS],
    ["Compression setting", 0.4],
    ["Actual measured token reduction (%)", 60.15],
    ["Pricing type", "Illustrative API-equivalent pricing"],
    ["Tokenizer", "tiktoken cl100k_base estimate"],
    ["Local Ollama API cost", "₹0 per API token"],
]

for row in assumption_rows:
    assumptions.append(row)

# Sheet 3: Interpretation
notes = workbook.create_sheet("Notes")
note_rows = [
    ["Cost Audit Notes"],
    ["Baseline uses the original 1,975-token RAG prompt."],
    ["Optimised uses the measured 787-token compressed prompt."],
    ["Output tokens are held constant at 200 for comparison."],
    ["API prices and INR conversion are illustrative assumptions."],
    ["Actual Ollama local inference has no per-token API charge."],
    ["Hardware electricity and operating costs are not included."],
    ["Token counts are estimates using cl100k_base."],
    ["LLMLingua reported a sequence-length warning; compression quality needs validation."],
]

for row in note_rows:
    notes.append(row)

# Basic formatting
for sheet in workbook.worksheets:
    sheet.freeze_panes = "A3" if sheet == ws else "A2"

    for cell in sheet[1]:
        cell.font = Font(bold=True, size=14)

    for cell in sheet[2] if sheet == ws else sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="365F91")

    for column in sheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )
        sheet.column_dimensions[
            get_column_letter(column[0].column)
        ].width = min(max_length + 4, 65)

# Number formats
for row in range(6, 10):
    for col in range(2, 5):
        ws.cell(row, col).number_format = '$0.000000'

for col in range(2, 5):
    ws.cell(9, col).number_format = '₹0.000000'

for col in range(2, 5):
    ws.cell(10, col).number_format = '0.00%'
    ws.cell(11, col).number_format = '0.00%'

file_path = output_dir / "cost_audit.xlsx"
workbook.save(file_path)

print("Cost audit workbook created successfully!")
print(f"Saved to: {file_path}")
print("Sheets: Cost Comparison, Assumptions, Notes")